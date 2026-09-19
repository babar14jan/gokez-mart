import os, requests
from io import BytesIO
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

API = "https://v2api.muktifresh.in/product/getproductbycatslug/client"
HEADERS = {"Origin": "https://v2.muktifresh.in"}
IMG_DIR = "muktifresh_images"
OUT_FILE = "muktifresh_organic_vegetables.xlsx"

os.makedirs(IMG_DIR, exist_ok=True)

data = requests.get(API, params={"slug": "organic-vegetables", "page": 1, "limit": 200}, headers=HEADERS).json()

wb = Workbook()
ws = wb.active
ws.title = "Organic Vegetables"
ws.append(["#", "Product Name", "Variant", "MRP (Rs)", "Selling Price (Rs)", "Image"])
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 45
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 12
ws.column_dimensions["E"].width = 18
ws.column_dimensions["F"].width = 20

row = 2
for i, p in enumerate(data, 1):
    name = p.get("display_name", "")
    img_url = p.get("featured_image") or p.get("product_image", "")
    attrs = p.get("attributes", [])

    safe_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in name)[:60]
    img_path = os.path.join(IMG_DIR, f"{safe_name}.png")

    if img_url and not os.path.exists(img_path):
        try:
            r = requests.get(img_url, timeout=10)
            PILImage.open(BytesIO(r.content)).convert("RGB").save(img_path, "PNG")
        except Exception as e:
            print(f"  Image failed: {name} — {e}")
            img_path = None

    def add_image(r):
        if img_path and os.path.exists(img_path):
            try:
                img = XLImage(img_path)
                img.width, img.height = 70, 55
                ws.add_image(img, f"F{r}")
            except Exception:
                ws.cell(r, 6, img_url)

    if not attrs:
        ws.cell(row, 1, i)
        ws.cell(row, 2, name)
        ws.cell(row, 3, "—")
        ws.cell(row, 4, p.get("mrp", 0))
        ws.cell(row, 5, p.get("sp", 0))
        ws.row_dimensions[row].height = 60
        add_image(row)
        row += 1
    else:
        for j, attr in enumerate(attrs):
            ws.cell(row, 1, i if j == 0 else "")
            ws.cell(row, 2, name if j == 0 else "")
            ws.cell(row, 3, attr.get("variance", ""))
            ws.cell(row, 4, attr.get("mrp", 0))
            ws.cell(row, 5, attr.get("sp", 0))
            ws.row_dimensions[row].height = 60
            if j == 0:
                add_image(row)
            row += 1

    print(f"[{i:02d}] {name}")

wb.save(OUT_FILE)
print(f"\nSaved: {OUT_FILE}  |  Images: {IMG_DIR}/  |  Products: {len(data)}")
